#!/usr/bin/env python3
"""把 Excel 某一列的“步骤详情”拆分为独立耗时列。

默认按用户指定的 AB 列拆分，并在 AB 后插入结果列。拆分规则与示例工作簿
Z 列后的列一致：

1. Model(...s) → 调用 Skill
2. Model(...s) → 调用 Workflow
3. call_nego_plan_llm(...s)
4. call_exec_llm(...s)
5. Model(...s) → 生成回答

脚本不会覆盖输入文件；除非显式传入 --overwrite，否则也不会覆盖已有输出文件。
"""

from __future__ import annotations

import argparse
import re
import sys
from copy import copy
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.utils.cell import range_boundaries


DEFAULT_INPUT = Path(
    "/Users/zhangfan/project/aico_timedelay_report/e2e_delay/4_9.4/"
    "查数_150_匹配all_operational_logs_4_9.4_步骤耗时拆分.xlsx"
)

CATEGORY_ORDER = ("skill", "workflow", "nego", "exec", "answer")

# 至少生成与 Z 列示例相同数量的列；若实际数据出现更多次数，会自动扩展。
REFERENCE_MINIMUMS = {
    "skill": 5,
    "workflow": 5,
    "nego": 5,
    "exec": 14,
    "answer": 5,
}

HEADER_TEMPLATES = {
    "skill": "调用Skill_{n}耗时(s)",
    "workflow": "调用Workflow_{n}耗时(s)",
    "nego": "call_nego_plan_llm_{n}耗时(s)",
    "exec": "call_exec_llm_{n}耗时(s)",
    "answer": "生成回答_{n}耗时(s)",
}

MODEL_LINE_RE = re.compile(
    r"^\s*Model\(\s*(?P<value>[-+]?\d+(?:\.\d+)?)\s*"
    r"(?P<unit>ms|s)\s*\)\s*(?:→|->|=>)\s*(?P<action>.+?)\s*$",
    re.IGNORECASE,
)

DIRECT_LINE_RE = re.compile(
    r"^\s*(?P<name>call_nego_plan_llm|call_exec_llm)"
    r"\(\s*(?P<value>[-+]?\d+(?:\.\d+)?)\s*"
    r"(?P<unit>ms|s)\s*\)\s*$",
    re.IGNORECASE,
)

GENERATED_HEADER_RE = re.compile(
    r"^(?:调用Skill_\d+耗时\(s\)|调用Workflow_\d+耗时\(s\)|"
    r"call_nego_plan_llm_\d+耗时\(s\)|call_exec_llm_\d+耗时\(s\)|"
    r"生成回答_\d+耗时\(s\))$"
)


@dataclass(frozen=True)
class SheetAnalysis:
    sheet_name: str
    nonempty_cells: int
    rows_with_matches: int
    total_matches: int
    maximums: dict[str, int]


def duration_in_seconds(value: str, unit: str) -> float:
    seconds = float(value)
    if unit.lower() == "ms":
        seconds /= 1000.0
    return seconds


def parse_step_details(value: object) -> dict[str, list[float]]:
    """提取一个单元格中的目标耗时，保留每类步骤在文本里的出现顺序。"""
    result: dict[str, list[float]] = {name: [] for name in CATEGORY_ORDER}
    if not isinstance(value, str) or not value.strip():
        return result

    for raw_line in value.splitlines():
        line = raw_line.strip()
        if not line:
            continue

        model_match = MODEL_LINE_RE.match(line)
        if model_match:
            action = model_match.group("action").strip()
            duration = duration_in_seconds(
                model_match.group("value"), model_match.group("unit")
            )
            if action.startswith("调用") and "skill" in action.lower():
                result["skill"].append(duration)
            elif action.startswith("调用") and "workflow" in action.lower():
                result["workflow"].append(duration)
            elif "生成回答" in action:
                result["answer"].append(duration)
            continue

        direct_match = DIRECT_LINE_RE.match(line)
        if direct_match:
            category = "nego" if direct_match.group("name").lower() == "call_nego_plan_llm" else "exec"
            result[category].append(
                duration_in_seconds(
                    direct_match.group("value"), direct_match.group("unit")
                )
            )

    return result


def iter_source_values(ws, source_col: int, header_row: int) -> Iterable[tuple[int, object]]:
    for row_number in range(header_row + 1, ws.max_row + 1):
        yield row_number, ws.cell(row=row_number, column=source_col).value


def analyze_sheet(ws, source_col: int, header_row: int) -> SheetAnalysis:
    maximums = {name: 0 for name in CATEGORY_ORDER}
    nonempty_cells = 0
    rows_with_matches = 0
    total_matches = 0

    for _, value in iter_source_values(ws, source_col, header_row):
        if value not in (None, ""):
            nonempty_cells += 1
        parsed = parse_step_details(value)
        row_match_count = sum(len(values) for values in parsed.values())
        if row_match_count:
            rows_with_matches += 1
            total_matches += row_match_count
        for category, values in parsed.items():
            maximums[category] = max(maximums[category], len(values))

    return SheetAnalysis(
        sheet_name=ws.title,
        nonempty_cells=nonempty_cells,
        rows_with_matches=rows_with_matches,
        total_matches=total_matches,
        maximums=maximums,
    )


def choose_worksheet(workbook, requested_name: str | None, source_col: int, header_row: int):
    if requested_name:
        if requested_name not in workbook.sheetnames:
            available = "、".join(workbook.sheetnames)
            raise ValueError(f"找不到工作表 {requested_name!r}。现有工作表：{available}")
        ws = workbook[requested_name]
        analysis = analyze_sheet(ws, source_col, header_row)
        return ws, analysis

    analyses = [analyze_sheet(ws, source_col, header_row) for ws in workbook.worksheets]
    candidates = [item for item in analyses if item.rows_with_matches > 0]
    if len(candidates) == 1:
        analysis = candidates[0]
        return workbook[analysis.sheet_name], analysis
    if not candidates:
        details = "；".join(
            f"{item.sheet_name}: 非空{item.nonempty_cells}格、匹配0行" for item in analyses
        )
        raise ValueError(
            f"没有在 {get_column_letter(source_col)} 列识别到步骤详情。"
            f"请确认列号，或用 --source-column 指定正确列。检查结果：{details}"
        )
    names = "、".join(item.sheet_name for item in candidates)
    raise ValueError(f"多个工作表的目标列都能匹配：{names}。请用 --sheet 明确指定。")


def make_layout(maximums: dict[str, int], compact: bool) -> list[tuple[str, int, str]]:
    layout: list[tuple[str, int, str]] = []
    for category in CATEGORY_ORDER:
        count = maximums[category]
        if not compact:
            count = max(count, REFERENCE_MINIMUMS[category])
        for occurrence in range(1, count + 1):
            layout.append(
                (
                    category,
                    occurrence,
                    HEADER_TEMPLATES[category].format(n=occurrence),
                )
            )
    return layout


def shifted_range(ref: str, insertion_col: int, amount: int) -> str:
    """把普通单区域 A1 引用按插列位置右移/扩展。"""
    min_col, min_row, max_col, max_row = range_boundaries(ref)
    if min_col >= insertion_col:
        min_col += amount
        max_col += amount
    elif max_col >= insertion_col:
        max_col += amount
    return (
        f"{get_column_letter(min_col)}{min_row}:"
        f"{get_column_letter(max_col)}{max_row}"
    )


def snapshot_and_shift_column_dimensions(ws, insertion_col: int, amount: int) -> None:
    snapshots = []
    for key, dimension in list(ws.column_dimensions.items()):
        col_idx = column_index_from_string(key)
        if col_idx >= insertion_col:
            snapshots.append((key, col_idx + amount, copy(dimension)))

    for old_key, _, _ in snapshots:
        del ws.column_dimensions[old_key]
    for _, new_idx, dimension in snapshots:
        new_key = get_column_letter(new_idx)
        dimension.index = new_key
        ws.column_dimensions[new_key] = dimension


def write_split_columns(ws, source_col: int, header_row: int, layout) -> None:
    insertion_col = source_col + 1
    amount = len(layout)

    next_header = ws.cell(row=header_row, column=insertion_col).value
    if isinstance(next_header, str) and GENERATED_HEADER_RE.match(next_header):
        raise ValueError(
            f"{ws.title}!{get_column_letter(insertion_col)}{header_row} 已经是拆分列 "
            f"{next_header!r}，为避免重复插入，脚本已停止。"
        )

    auto_filter_ref = ws.auto_filter.ref
    table_refs = {table.name: table.ref for table in ws.tables.values()}

    snapshot_and_shift_column_dimensions(ws, insertion_col, amount)
    ws.insert_cols(insertion_col, amount)

    # openpyxl 插列时不会自动更新这些区域引用，因此在此同步调整。
    if auto_filter_ref:
        ws.auto_filter.ref = shifted_range(auto_filter_ref, insertion_col, amount)
    for table in ws.tables.values():
        old_ref = table_refs.get(table.name)
        if old_ref:
            table.ref = shifted_range(old_ref, insertion_col, amount)

    source_header = ws.cell(row=header_row, column=source_col)
    for offset, (_, _, header) in enumerate(layout):
        col_idx = insertion_col + offset
        header_cell = ws.cell(row=header_row, column=col_idx, value=header)
        header_cell._style = copy(source_header._style)
        header_cell.font = copy(source_header.font)
        header_cell.fill = copy(source_header.fill)
        header_cell.border = copy(source_header.border)
        header_cell.alignment = copy(source_header.alignment)
        header_cell.protection = copy(source_header.protection)
        ws.column_dimensions[get_column_letter(col_idx)].width = 16

    for row_number, source_value in iter_source_values(ws, source_col, header_row):
        parsed = parse_step_details(source_value)
        source_cell = ws.cell(row=row_number, column=source_col)
        for offset, (category, occurrence, _) in enumerate(layout):
            target = ws.cell(row=row_number, column=insertion_col + offset)
            target._style = copy(source_cell._style)
            target.font = copy(source_cell.font)
            target.fill = copy(source_cell.fill)
            target.border = copy(source_cell.border)
            target.protection = copy(source_cell.protection)
            alignment = copy(source_cell.alignment)
            alignment.wrap_text = False
            target.alignment = alignment
            target.number_format = "0.000"
            values = parsed[category]
            if occurrence <= len(values):
                target.value = values[occurrence - 1]


def default_output_path(input_path: Path, source_column: str) -> Path:
    return input_path.with_name(
        f"{input_path.stem}_{source_column}步骤耗时拆分{input_path.suffix}"
    )


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="将 Excel 指定列中的步骤详情拆分成独立耗时列。"
    )
    parser.add_argument(
        "input",
        nargs="?",
        type=Path,
        default=DEFAULT_INPUT,
        help=f"输入工作簿，默认：{DEFAULT_INPUT}",
    )
    parser.add_argument("--output", "-o", type=Path, help="输出工作簿路径")
    parser.add_argument("--sheet", help="目标工作表名称；不填时自动识别")
    parser.add_argument(
        "--source-column",
        default="AB",
        help="步骤详情所在列，支持 Excel 列字母或正整数，默认 AB",
    )
    parser.add_argument("--header-row", type=int, default=1, help="表头行号，默认 1")
    parser.add_argument(
        "--compact",
        action="store_true",
        help="只生成数据实际需要的列，不补齐到 Z 列示例的固定宽度",
    )
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="只检查匹配结果，不写出工作簿",
    )
    parser.add_argument(
        "--overwrite",
        action="store_true",
        help="允许覆盖已存在的输出文件（仍不会覆盖输入文件）",
    )
    return parser.parse_args()


def parse_column(value: str) -> int:
    text = value.strip().upper()
    if text.isdigit():
        column = int(text)
        if column < 1:
            raise ValueError("列号必须大于 0")
        return column
    try:
        return column_index_from_string(text)
    except ValueError as exc:
        raise ValueError(f"无效列号：{value!r}") from exc


def main() -> int:
    args = parse_args()
    input_path = args.input.expanduser().resolve()
    source_col = parse_column(args.source_column)
    source_column_name = get_column_letter(source_col)
    output_path = (
        args.output or default_output_path(input_path, source_column_name)
    ).expanduser().resolve()

    if not input_path.is_file():
        raise FileNotFoundError(f"输入文件不存在：{input_path}")
    if input_path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError("仅支持 .xlsx 或 .xlsm 文件")
    if args.header_row < 1:
        raise ValueError("--header-row 必须大于 0")
    if output_path == input_path:
        raise ValueError("输出路径不能与输入文件相同，请保留原文件并指定新的输出路径")
    if output_path.exists() and not args.overwrite and not args.dry_run:
        raise FileExistsError(f"输出文件已存在：{output_path}；如需覆盖请加 --overwrite")

    keep_vba = input_path.suffix.lower() == ".xlsm"
    workbook = load_workbook(input_path, data_only=False, keep_vba=keep_vba)
    ws, analysis = choose_worksheet(workbook, args.sheet, source_col, args.header_row)
    layout = make_layout(analysis.maximums, args.compact)

    print(f"工作表：{analysis.sheet_name}")
    print(f"源列：{get_column_letter(source_col)}")
    print(f"非空单元格：{analysis.nonempty_cells}")
    print(f"成功匹配行数：{analysis.rows_with_matches}")
    print(f"匹配耗时总数：{analysis.total_matches}")
    print(
        "每行最大出现次数："
        + "，".join(f"{name}={analysis.maximums[name]}" for name in CATEGORY_ORDER)
    )
    print(f"将生成列数：{len(layout)}")

    if analysis.rows_with_matches == 0:
        raise ValueError("目标列没有可拆分内容，未生成文件")
    if args.dry_run:
        print("dry-run 完成，未写入文件。")
        return 0

    write_split_columns(ws, source_col, args.header_row, layout)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    print(f"已生成：{output_path}")
    return 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    except Exception as exc:
        print(f"错误：{exc}", file=sys.stderr)
        sys.exit(1)
