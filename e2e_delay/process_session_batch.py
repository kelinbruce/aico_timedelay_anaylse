#!/usr/bin/env python3
"""解压日志 → 调用 analyze → 按问数表 Session ID 匹配。Python 3.10+ / openpyxl。"""
from __future__ import annotations

import argparse
import csv
import gzip
import hashlib
import json
import re
import subprocess
import sys
import tarfile
import zipfile
from collections import Counter
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

# 内置历史拆分规则，入口脚本可独立复制，无需仓库辅助模块。
CATEGORY_ORDER = ("skill", "workflow", "nego", "exec", "answer")

# 至少生成与 Z 列示例相同数量的列；若实际数据出现更多次数，会自动扩展。
REFERENCE_MINIMUMS = {
    "skill": 5,
    "workflow": 5,
    "nego": 5,
    "exec": 14,
    "answer": 5,
}

HEADER_TEMPLATES = {
    "skill": "调用Skill_{n}耗时(s)",
    "workflow": "调用Workflow_{n}耗时(s)",
    "nego": "call_nego_plan_llm_{n}耗时(s)",
    "exec": "call_exec_llm_{n}耗时(s)",
    "answer": "生成回答_{n}耗时(s)",
}

MODEL_LINE_RE = re.compile(
    r"^\s*Model\(\s*(?P<value>[-+]?\d+(?:\.\d+)?)\s*"
    r"(?P<unit>ms|s)\s*\)\s*(?:→|->|=>)\s*(?P<action>.+?)\s*$",
    re.IGNORECASE,
)

DIRECT_LINE_RE = re.compile(
    r"^\s*(?P<name>call_nego_plan_llm|call_exec_llm)"
    r"\(\s*(?P<value>[-+]?\d+(?:\.\d+)?)\s*"
    r"(?P<unit>ms|s)\s*\)\s*$",
    re.IGNORECASE,
)

def duration_in_seconds(value: str, unit: str) -> float:
    seconds = float(value)
    if unit.lower() == "ms":
        seconds /= 1000.0
    return seconds


def parse_step_details(value: object) -> dict[str, list[float]]:
    """提取一个单元格中的目标耗时，保留每类步骤在文本里的出现顺序。"""
    result: dict[str, list[float]] = {name: [] for name in CATEGORY_ORDER}
    if not isinstance(value, str) or not value.strip():
        return result

    for raw_line in value.splitlines():
        line = raw_line.strip()
        if not line:
            continue

        model_match = MODEL_LINE_RE.match(line)
        if model_match:
            action = model_match.group("action").strip()
            duration = duration_in_seconds(
                model_match.group("value"), model_match.group("unit")
            )
            if action.startswith("调用") and "skill" in action.lower():
                result["skill"].append(duration)
            elif action.startswith("调用") and "workflow" in action.lower():
                result["workflow"].append(duration)
            elif "生成回答" in action:
                result["answer"].append(duration)
            continue

        direct_match = DIRECT_LINE_RE.match(line)
        if direct_match:
            category = "nego" if direct_match.group("name").lower() == "call_nego_plan_llm" else "exec"
            result[category].append(
                duration_in_seconds(
                    direct_match.group("value"), direct_match.group("unit")
                )
            )

    return result


def text(value):
    return '' if value is None else str(value).strip()


def read_questions(path, sheet=None, session_column=None):
    if path.suffix.lower() == '.csv':
        if sheet:
            raise ValueError('--sheet 仅用于 Excel')
        with path.open(encoding='utf-8-sig', newline='') as handle:
            matrix = list(csv.reader(handle))
    elif path.suffix.lower() in ('.xlsx', '.xlsm'):
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            ws = wb[sheet] if sheet else wb.worksheets[0]
            matrix = list(ws.values)
        finally:
            wb.close()
    else:
        raise ValueError('问数文件仅支持 .csv、.xlsx、.xlsm；旧版 .xls 请先另存为 .xlsx')
    if not matrix:
        raise ValueError('问数文件为空')
    headers = [text(v) for v in matrix[0]]
    candidates = [session_column] if session_column else [
        '额外答案-sessionId', 'sessionId', 'session_id', 'Session ID', 'SessionID', '会话ID',
    ]
    found = [h for h in candidates if h in headers]
    if len(found) != 1:
        raise ValueError(f'无法唯一确定Session列（找到{found}），请用 --session-column 指定')
    used = {found[0], '编号', 'QA_ID', '问题', '单步/多步'}
    if any(headers.count(h) > 1 for h in used):
        raise ValueError('问数文件的Session或题目信息列重名，请先消除歧义')
    rows = []
    for number, values in enumerate(matrix[1:], 2):
        if not any(v is not None and v != '' for v in values):
            continue
        rows.append((number, dict(zip(headers, values))))
    if not rows:
        raise ValueError('问数表没有数据行')
    return rows, found[0]


def archive_kind(path):
    name = path.name.lower()
    if name.endswith('.zip'):
        return 'zip'
    if name.endswith(('.tar', '.tar.gz', '.tgz', '.tar.bz2', '.tbz2', '.tar.xz', '.txz')):
        return 'tar'
    return None


class Extractor:
    """仅提取普通文件，阻止越界、链接、重名覆盖及无限嵌套。"""
    def __init__(self, max_bytes):
        self.max_bytes = max_bytes
        self.bytes = 0

    def copy(self, source, dest):
        dest.parent.mkdir(parents=True, exist_ok=True)
        with dest.open('xb') as output:
            while block := source.read(1024 * 1024):
                self.bytes += len(block)
                if self.bytes > self.max_bytes:
                    raise ValueError('解压总量超限，可用 --max-expanded-gb 调整')
                output.write(block)

    def extract(self, archive, root, depth=0):
        if depth > 5:
            raise ValueError('压缩包嵌套超过5层')
        root.mkdir(parents=True, exist_ok=False)
        files = []

        def target(name):
            name = name.replace('\\', '/')
            dest = (root / name).resolve()
            if not dest.is_relative_to(root.resolve()) or ':' in name:
                raise ValueError(f'压缩包路径非法: {name}')
            return dest

        if archive_kind(archive) == 'zip':
            with zipfile.ZipFile(archive) as z:
                for member in z.infolist():
                    dest = target(member.filename)
                    if (member.external_attr >> 16) & 0o170000 == 0o120000:
                        raise ValueError(f'压缩包包含符号链接: {member.filename}')
                    if member.is_dir():
                        dest.mkdir(parents=True, exist_ok=True)
                    else:
                        with z.open(member) as handle:
                            self.copy(handle, dest)
                        files.append(dest)
        elif archive_kind(archive) == 'tar':
            with tarfile.open(archive) as tar:
                for member in tar:
                    dest = target(member.name)
                    if member.isdir():
                        dest.mkdir(parents=True, exist_ok=True)
                    elif member.isfile():
                        with tar.extractfile(member) as handle:
                            self.copy(handle, dest)
                        files.append(dest)
                    else:
                        raise ValueError(f'压缩包包含非普通文件: {member.name}')
        else:
            raise ValueError(f'不支持的压缩包: {archive.name}')
        for number, path in enumerate(files):
            if archive_kind(path) and '__MACOSX' not in path.parts:
                self.extract(path, root.parent / f'{root.name}_nested_{number}', depth + 1)


def collect_logs(raw, output, extractor):
    output.mkdir()
    manifest, seen, session_ids = [], set(), set()
    sources = sorted(p for p in raw.rglob('*') if p.is_file()
                     and p.name.startswith('nextagent-operational')
                     and not archive_kind(p) and '__MACOSX' not in p.parts)
    if not sources:
        raise ValueError('压缩包内未找到 nextagent-operational 开头的日志文件')
    for number, source in enumerate(sources, 1):
        dest = output / f'nextagent-operational-{number:05d}.jsonl'
        counts = Counter()
        opener = gzip.open if source.name.endswith('.gz') else open
        with opener(source, 'rb') as handle, dest.open('xb') as target:
            for line in handle:
                extractor.bytes += len(line)
                if extractor.bytes > extractor.max_bytes:
                    raise ValueError('解压总量超限，可用 --max-expanded-gb 调整')
                stripped = line.strip()
                if not stripped:
                    continue
                digest = hashlib.sha256(stripped).digest()
                if digest in seen:
                    counts['duplicate_lines'] += 1
                    continue
                seen.add(digest)
                target.write(stripped + b'\n')
                counts['kept_lines'] += 1
                try:
                    event = json.loads(stripped)
                    if isinstance(event, dict) and event.get('sessionId'):
                        session_ids.add(text(event['sessionId']))
                except (ValueError, UnicodeDecodeError):
                    counts['non_json_lines'] += 1
        manifest.append({'source': str(source.relative_to(raw)), 'output': dest.name, **counts})
    if not any(item.get('kept_lines', 0) for item in manifest):
        raise ValueError('日志没有非空内容')
    return manifest, session_ids


def match_rows(questions, session_column, report, raw_ids):
    wb = load_workbook(report, read_only=True, data_only=True)
    try:
        matrix = list(wb['会话明细'].values)
    finally:
        wb.close()
    headers = [text(h) or f'原报告无表头列（第{i+1}列）' for i, h in enumerate(matrix[0])]
    if len(headers) != len(set(headers)):
        raise ValueError('分析报告含重复列名')
    sid_col = headers.index('会话ID')
    step_col = headers.index('步骤详情(每步操作)')
    index = {}
    for row in matrix[1:]:
        sid = text(row[sid_col])
        if not sid:
            continue
        if sid in index:
            raise ValueError(f'分析报告存在重复Session ID: {sid}')
        index[sid] = list(row)
    counts = Counter(text(src.get(session_column)) for _, src in questions)
    parsed = {sid: parse_step_details(row[step_col]) for sid, row in index.items() if sid in counts}
    limits = {k: max([REFERENCE_MINIMUMS[k]] + [len(p[k]) for p in parsed.values()]) for k in CATEGORY_ORDER}
    split_headers = [HEADER_TEMPLATES[k].format(n=i) for k in CATEGORY_ORDER for i in range(1, limits[k]+1)]
    extra = ['单步/多步', '来源行号', '查数编号', 'QA_ID', '查数问题', '匹配状态', '同Session查数记录数']
    final_headers = headers[:step_col+1] + split_headers + headers[step_col+1:] + extra
    if len(final_headers) != len(set(final_headers)):
        raise ValueError('分析报告已含拆分或匹配列，请指定未加工的 analyze 分析脚本')
    result = [final_headers]
    missing = [['来源行号', '查数编号', 'QA_ID', 'Session ID', '查数问题', '原因']]
    statuses = Counter()
    for number, src in questions:
        sid = text(src.get(session_column))
        if sid in index:
            raw = index[sid]
            split = [v for k in CATEGORY_ORDER for v in parsed[sid][k] + [None]*(limits[k]-len(parsed[sid][k]))]
            row = raw[:step_col+1] + split + raw[step_col+1:]
            status = '已匹配（共享会话）' if counts[sid] > 1 else '已匹配'
        else:
            row = [None] * (len(headers)+len(split_headers))
            row[sid_col] = sid or None
            status = ('源Session ID为空' if not sid else
                      '原始日志存在，未纳入分析报告' if sid in raw_ids else '报告中未匹配')
            missing.append([number, src.get('编号'), src.get('QA_ID'), sid or None, src.get('问题'), status])
        statuses[status] += 1
        result.append(row + [src.get('单步/多步'), number, src.get('编号'), src.get('QA_ID'),
                             src.get('问题'), status, counts[sid] if sid else None])
    summary = {'查数记录数': len(questions), '匹配记录数': sum(v for k, v in statuses.items() if k.startswith('已匹配')),
               '匹配唯一Session数': len(parsed), '空Session记录数': counts.get('', 0),
               '非空未匹配记录数': sum(v for k, v in counts.items() if k and k not in index),
               '重复Session组数': sum(v > 1 for k, v in counts.items() if k),
               '单步记录数': sum(src.get('单步/多步') == '单步' for _, src in questions),
               '多步记录数': sum(src.get('单步/多步') == '多步' for _, src in questions)}
    return result, missing, summary


def export_result(path, sheets):
    wb = Workbook()
    wb.remove(wb.active)
    for number, (name, rows) in enumerate(sheets, 1):
        ws = wb.create_sheet(name)
        for row in rows:
            ws.append(row)
        # 文本即文本，防止问题中的“=”被Excel当作公式。
        for row in ws:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = 's'
                cell.font = Font(name='Arial', size=10)
                cell.alignment = Alignment(vertical='center', wrap_text=True)
        ws.sheet_view.showGridLines = False
        ws.freeze_panes = 'B2' if name == '最终筛选结果' else 'A2'
        ws.row_dimensions[1].height = 48
        for cell in ws[1]:
            cell.fill = PatternFill('solid', fgColor='2F5496')
            cell.font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
        for col, header in enumerate(rows[0], 1):
            h = str(header)
            width = 18
            if 'ID' in h or h == 'Session ID':
                width = 48
            if any(k in h for k in ('步骤详情', '明细', '问题', '标题')):
                width = 65
            if '时间' in h:
                width = 24
            ws.column_dimensions[get_column_letter(col)].width = width
            if '耗时' in h and '明细' not in h:
                for cells in ws.iter_rows(min_row=2, min_col=col, max_col=col):
                    if isinstance(cells[0].value, (float, int)):
                        cells[0].number_format = '0.000'
        for i in range(2, len(rows)+1):
            ws.row_dimensions[i].height = 54
        if name == '匹配说明':
            ws.column_dimensions['A'].width = 28
            ws.column_dimensions['B'].width = 110
        elif len(rows) > 1:
            table = Table(displayName=f'ResultTable{number}', ref=ws.dimensions)
            table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
            ws.add_table(table)
    wb.save(path)
    wb.close()
    check = load_workbook(path, read_only=True, data_only=False)
    try:
        for name, expected in sheets:
            actual = list(check[name].values)
            normalized = [tuple(None if v == '' else v for v in row) for row in expected]
            if actual != normalized:
                raise ValueError(f'Excel导出核验失败: {name}')
    finally:
        check.close()


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--archive', type=Path, required=True, help='ZIP/TAR日志压缩包')
    parser.add_argument('--questions', type=Path, required=True, help='问数表 .xlsx/.xlsm/.csv，首行为表头')
    parser.add_argument('--output-dir', type=Path, required=True, help='新的输出目录，已存在时拒绝覆盖')
    parser.add_argument('--analyzer', type=Path, default=Path(__file__).resolve().parent/'1/analyze_65_kq_v3.py')
    parser.add_argument('--scenario', choices=['wireless', 'cloud'], default='wireless')
    parser.add_argument('--sheet', help='Excel工作表名，默认第一张')
    parser.add_argument('--session-column', help='问数表的Session列名，默认自动识别')
    parser.add_argument('--t-start', help='传给analyze的开始时间，UTC: YYYY-MM-DD HH:MM:SS')
    parser.add_argument('--t-end', help='传给analyze的结束时间，UTC: YYYY-MM-DD HH:MM:SS')
    parser.add_argument('--max-expanded-gb', type=float, default=10, help='所有解压层及规范化日志的累计大小上限，默认10 GiB')
    args = parser.parse_args(argv)
    for key in ('archive', 'questions', 'analyzer'):
        value = getattr(args, key).resolve()
        if not value.is_file():
            parser.error(f'文件不存在: {value}')
        setattr(args, key, value)
    if bool(args.t_start) != bool(args.t_end):
        parser.error('--t-start 和 --t-end 必须一起提供')
    if not 0 < args.max_expanded_gb < float('inf'):
        parser.error('--max-expanded-gb 必须为有限正数')
    if not archive_kind(args.archive):
        parser.error('仅支持 ZIP/TAR/TAR.GZ/TGZ/TAR.BZ2/TAR.XZ 压缩包')
    questions, session_column = read_questions(args.questions, args.sheet, args.session_column)
    out = args.output_dir.resolve()
    out.mkdir(parents=True, exist_ok=False)
    print('[1/4] 解压及汇集日志', flush=True)
    extractor = Extractor(int(args.max_expanded_gb * 1024**3))
    raw = out/'原始日志'
    raw.mkdir()
    extractor.extract(args.archive, raw/'archive')
    logs = out/'all_operational_logs'
    manifest, raw_ids = collect_logs(raw, logs, extractor)
    (out/'解压文件清单.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding='utf-8')
    print(f'[2/4] 调用分析脚本，日志文件数: {len(manifest)}', flush=True)
    command = [sys.executable, str(args.analyzer), str(logs), args.scenario]
    if args.t_start:
        command += ['--t-start', args.t_start, '--t-end', args.t_end]
    with (out/'解析运行记录.log').open('w', encoding='utf-8') as log:
        completed = subprocess.run(command, cwd=out, stdout=log, stderr=subprocess.STDOUT, check=False)
    if completed.returncode:
        raise RuntimeError(f'analyze执行失败（退出码{completed.returncode}），详见 {out / "解析运行记录.log"}')
    report = out/f'all_operational_logs_knowledge-qa性能分析_{"无线" if args.scenario == "wireless" else "云核"}.xlsx'
    if not report.is_file():
        raise ValueError(f'analyze未生成预期Excel: {report}；请选择直接输出XLSX的原始分析脚本')
    print('[3/4] 按Session ID精确匹配并拆分步骤耗时', flush=True)
    result, missing, summary = match_rows(questions, session_column, report, raw_ids)
    notes = [['项目', '内容'], *[[k, v] for k, v in summary.items()],
             ['问数来源', f'{args.questions.name}；工作表：{args.sheet or "默认第一张（CSV不适用）"}；Session列：{session_column}'],
             ['日志来源', args.archive.name], ['分析脚本', args.analyzer.name], ['会话来源', report.name+' / 会话明细'],
             ['匹配规则', 'Session去首尾空白后精确匹配；保留问数顺序、重复Session和缺失记录。不按问题文本补配。'],
             ['重复Session口径', '共享Session引用同一会话耗时，按会话统计须先去重。'],
             ['步骤耗时口径', '调用Skill/Workflow为发起该调用的Model耗时；nego/exec为对应阶段耗时；同类按出现顺序拆列，单位秒。'],
             ['单步/多步来源', '直接取问数表同名列；源列不存在则留空，不依据日志推断。'],
             ['来源行号', 'Excel为原工作表行号；CSV为逻辑记录号（包含表头），跳过整行空白。'],
             ['日志去重', '去除首尾空白后按整行SHA256跨文件去重，不按Session去重。'],
             ['非JSON日志行数', sum(item.get('non_json_lines', 0) for item in manifest)],
             ['原始日志存在标记', '仅依据顶层JSON sessionId确认存在；其他日志格式仍由analyze处理。'],
             ['分析口径', '沿用指定analyze脚本的会话筛选、时间和耗时计算规则；详情见解析运行记录。']]
    target = out/f'最终SessionID筛选结果_{len(questions)}行_步骤耗时拆分_单步多步标记.xlsx'
    print('[4/4] 导出并重新读取核验', flush=True)
    export_result(target, [('最终筛选结果', result), ('匹配说明', notes), ('未匹配SessionID', missing)])
    (out/'匹配汇总.json').write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    print(f'完成: {target}')
    return 0


if __name__ == '__main__':
    try:
        sys.exit(main())
    except (OSError, ValueError, KeyError, RuntimeError, zipfile.BadZipFile, tarfile.TarError) as exc:
        print(f'处理失败: {exc}', file=sys.stderr)
        sys.exit(1)
