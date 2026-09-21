"""运行: python -m unittest discover -s e2e_delay -p test_process_session_batch.py"""
import gzip
import io
import shutil
import subprocess
import sys
import tarfile
import tempfile
import unittest
import zipfile
from pathlib import Path

from openpyxl import Workbook, load_workbook
from process_session_batch import Extractor, collect_logs, export_result, match_rows, read_questions


class BatchTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def test_standalone_script_without_repository_modules(self):
        script = self.root/'process_session_batch.py'
        shutil.copyfile(Path(__file__).with_name(script.name), script)
        # -I排除PYTHONPATH、当前目录，验证不再依赖仓库辅助模块。
        result = subprocess.run([sys.executable, '-I', str(script), '--help'],
                                cwd=self.root, capture_output=True, text=True)
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertIn('--questions', result.stdout)
        result = subprocess.run([sys.executable, '-I', str(script), '--archive', 'logs.zip'],
                                cwd=self.root, capture_output=True, text=True)
        self.assertEqual(result.returncode, 2)
        self.assertIn('--questions', result.stderr)
        self.assertIn('--output-dir', result.stderr)

    def test_nested_archives_rotated_gzip_and_dedup(self):
        line = b'{"sessionId":"s1"}\n'
        nested = io.BytesIO()
        with zipfile.ZipFile(nested, 'w') as z:
            z.writestr('nextagent-operational.jsonl.1.gz', gzip.compress(line))
        archive = self.root/'logs.zip'
        with zipfile.ZipFile(archive, 'w') as z:
            z.writestr('nested.zip', nested.getvalue())
            z.writestr('dir/nextagent-operational.jsonl', line + b'{"sessionId":"s2"}\n')
        raw = self.root/'raw'
        raw.mkdir()
        extractor = Extractor(100000)
        extractor.extract(archive, raw/'archive')
        manifest, ids = collect_logs(raw, self.root/'logs', extractor)
        self.assertEqual(ids, {'s1', 's2'})
        self.assertEqual(sum(x.get('duplicate_lines', 0) for x in manifest), 1)
        self.assertEqual(sum(x.get('kept_lines', 0) for x in manifest), 2)

    def test_traversal_and_expansion_limit(self):
        archive = self.root/'bad.zip'
        with zipfile.ZipFile(archive, 'w') as z:
            z.writestr('../escaped.txt', 'bad')
        with self.assertRaises(ValueError):
            Extractor(1000).extract(archive, self.root/'raw')
        self.assertFalse((self.root/'escaped.txt').exists())
        archive = self.root/'big.zip'
        with zipfile.ZipFile(archive, 'w') as z:
            z.writestr('file', '123456')
        with self.assertRaises(ValueError):
            Extractor(5).extract(archive, self.root/'limited')

    def test_tar_and_link_rejection(self):
        path = self.root/'logs.tgz'
        with tarfile.open(path, 'w:gz') as tar:
            info = tarfile.TarInfo('nextagent-operational.jsonl')
            info.size = 3
            tar.addfile(info, io.BytesIO(b'{}\n'))
        Extractor(1000).extract(path, self.root/'tar')
        self.assertEqual((self.root/'tar/nextagent-operational.jsonl').read_bytes(), b'{}\n')
        with tarfile.open(self.root/'link.tar', 'w') as tar:
            info = tarfile.TarInfo('link')
            info.type = tarfile.SYMTYPE
            info.linkname = '/tmp'
            tar.addfile(info)
        with self.assertRaises(ValueError):
            Extractor(1000).extract(self.root/'link.tar', self.root/'links')

    def test_excel_matching_missing_shared_sessions_and_literal_text(self):
        wb = Workbook()
        ws = wb.active
        ws.title = '问数'
        ws.append(['sessionId', '编号', '问题', '单步/多步'])
        for row in [[' s1 ', 1, '第一题', '单步'], ['s1', 2, '第二题', '多步'],
                    [None, 3, '缺失', '单步'], ['excluded', 4, '排除', '单步']]:
            ws.append(row)
        questions_file = self.root/'questions.xlsx'
        wb.save(questions_file)
        questions, sid_col = read_questions(questions_file, '问数')
        questions[0][1]['问题'] = '=1+1'
        report = Workbook()
        ws = report.active
        ws.title = '会话明细'
        ws.append(['会话ID', '会话总耗时(秒)', '步骤详情(每步操作)'])
        ws.append(['s1', 2.5, 'Model(100ms) → 调用 Skill\ncall_exec_llm(2s)'])
        report_file = self.root/'report.xlsx'
        report.save(report_file)
        rows, missing, summary = match_rows(questions, sid_col, report_file, {'excluded'})
        self.assertEqual(summary['匹配记录数'], 2)
        self.assertEqual(summary['匹配唯一Session数'], 1)
        self.assertEqual(summary['空Session记录数'], 1)
        self.assertEqual(summary['非空未匹配记录数'], 1)
        self.assertIn('原始日志存在', missing[-1][-1])
        self.assertEqual(rows[1][rows[0].index('调用Skill_1耗时(s)')], 0.1)
        output = self.root/'result.xlsx'
        export_result(output, [('最终筛选结果', rows), ('未匹配SessionID', missing)])
        saved = load_workbook(output)
        cell = saved['最终筛选结果'].cell(2, rows[0].index('查数问题')+1)
        self.assertEqual((cell.value, cell.data_type), ('=1+1', 's'))
        saved.close()

    def test_ambiguous_session_columns(self):
        path = self.root/'q.csv'
        path.write_text('sessionId,会话ID\na,b\n')
        with self.assertRaises(ValueError):
            read_questions(path)
        rows, column = read_questions(path, session_column='会话ID')
        self.assertEqual((rows[0][1][column], column), ('b', '会话ID'))


if __name__ == '__main__':
    unittest.main()
